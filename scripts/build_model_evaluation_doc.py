from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "documentation" / "OneEye_Model_Evaluation_and_System_Verification.docx"

NAVY = "12324A"
PALE_BLUE = "EDF7FB"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
DARK_GRAY = RGBColor(68, 77, 86)


def set_cell_fill(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_margins(cell, top: int = 110, start: int = 120, bottom: int = 110, end: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), BORDER)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def set_run_font(run, size: float, *, bold: bool = False, color: RGBColor = BLACK) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_body(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, 10.5, bold=True)
        rest = paragraph.add_run(text[len(bold_lead) :])
        set_run_font(rest, 10.5, color=DARK_GRAY)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, 10.5, color=DARK_GRAY)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    set_run_font(run, 10.2, color=DARK_GRAY)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    table.rows[0].cells[0].paragraphs[0].paragraph_format.keep_with_next = True
    set_repeat_table_header(table.rows[0])
    for index, (cell, header, width) in enumerate(zip(table.rows[0].cells, headers, widths, strict=True)):
        cell.width = Inches(width)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell)
        set_cell_borders(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        set_run_font(run, 9.2, bold=True, color=RGBColor(255, 255, 255))
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        fill = PALE_BLUE if row_index % 2 == 0 else "FFFFFF"
        for index, (cell, value, width) in enumerate(zip(cells, values, widths, strict=True)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_fill(cell, fill)
            set_cell_margins(cell)
            set_cell_borders(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index in {0, 1} else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            set_run_font(run, 8.8, bold=(index == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK_GRAY
    for name, size, before, after in (
        ("Title", 25, 0, 10),
        ("Heading 1", 17, 15, 7),
        ("Heading 2", 12.5, 11, 5),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    title_properties = styles["Title"]._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)


def add_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("OneEye Model Evaluation and System Verification  |  13 September 2026")
        set_run_font(run, 8.2, color=RGBColor(105, 113, 120))


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    configure_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run("OneEye Model Evaluation and System Verification")
    set_run_font(title_run, 25, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("Final evidence report for technical review and judging")
    set_run_font(run, 12.5, bold=True, color=RGBColor(19, 122, 157))

    add_body(
        doc,
        "OneEye is a five-event surveillance system that combines a browser dashboard, FastAPI services, "
        "SQLite event storage, OpenCV camera capture, and Ultralytics inference. The implementation and guarded "
        "training sequence are complete. Every reported result comes from a held-out test split, and a model is "
        "promoted only when every required score is strictly greater than 95 percent.",
    )
    add_body(
        doc,
        "Final decision. The expanded fallen-person classifier passed the promotion gate and is deployed. Fight, "
        "garbage, expanded mobile-phone, and accident candidates did not meet every required score, so their "
        "research checkpoints were preserved and were not presented as production-grade models.",
        bold_lead="Final decision.",
    )

    doc.add_heading("Model Outcome Summary", level=1)
    add_table(
        doc,
        ["Model", "Untouched test result", "Strict gate", "Runtime decision"],
        [
            ["Fight", "Acc 83.75%  P 91.41%  R 74.50%", "Failed", "Baseline retained"],
            ["Garbage", "P 67.54%  R 44.49%  mAP50 49.04%", "Failed", "Baseline retained"],
            ["Fallen Person", "Acc 100%  P 100%  R 100%", "Passed", "Expanded model exported"],
            ["Mobile Phone", "P 87.31%  R 85.10%  mAP50 89.22%", "Failed", "Baseline retained"],
            ["Accident", "Acc 79.42%  P 79.04%  R 80.09%", "Failed", "Baseline retained"],
        ],
        [1.25, 3.05, 0.85, 1.75],
    )
    add_body(
        doc,
        "The gate is intentionally strict. Confidence thresholds between 91 and 97 percent control whether an "
        "individual alert is emitted; they do not turn a model with lower held-out accuracy into a 95 percent "
        "accurate model.",
    )

    doc.add_page_break()
    doc.add_heading("Dataset and Training Evidence", level=1)
    add_table(
        doc,
        ["Model", "Prepared data", "Separation rule", "Training result"],
        [
            [
                "Fight",
                "4,800 four-frame mosaics from 2,000 RWF-2000 videos",
                "Official 1,600 development and 400 test "
                "video boundary",
                "R3D-18 and multiple-instance research completed",
            ],
            [
                "Garbage",
                "2,527 train  225 validation  225 test",
                "TACO validation and test held fixed; "
                "MJU training only",
                "YOLO11s completed 100 epochs",
            ],
            [
                "Fallen Person",
                "937 train  185 validation  163 test frames",
                "Video groups kept disjoint; ambiguous "
                "transition frames excluded",
                "YOLO11s classifier stopped at epoch 16",
            ],
            [
                "Mobile Phone",
                "2,166 train  538 validation  793 test",
                "Open Images source boundaries preserved",
                "YOLO11s detector stopped at epoch 40; "
                "best epoch 25",
            ],
            [
                "Accident",
                "2,100 train  448 validation  452 test mosaics",
                "Each of 1,500 Nexar videos belongs to one "
                "class and one split",
                "YOLO11s classifier stopped at epoch 32; "
                "best epoch 17",
            ],
        ],
        [1.15, 1.85, 2.15, 1.75],
    )
    doc.add_heading("Data Integrity Controls", level=2)
    for item in (
        "All five dataset validators completed with zero reported errors.",
        "Detection validators check image and label pairs, normalized geometry, "
        "corruption, missing files, and split structure.",
        "Temporal validators check exact samples per source video, ordered frame "
        "timestamps, class balance, manifest agreement, and cross-split leakage.",
        "Untouched test sets were evaluated only after training completed and the "
        "best validation checkpoint had been selected.",
    ):
        add_bullet(doc, item)

    doc.add_heading("Training Controls", level=2)
    add_body(
        doc,
        "Training used deterministic seeds, CUDA on an NVIDIA GeForce RTX 5050 Laptop GPU, per-epoch last.pt "
        "checkpoints, validation-selected best.pt checkpoints, and atomic JSON status updates. A guarded queue "
        "accepted training as complete only after the pipeline wrote an explicit completion marker.",
    )

    doc.add_page_break()
    doc.add_heading("Detailed Held Out Results", level=1)
    doc.add_heading("Fight Detection", level=2)
    add_body(
        doc,
        "The strongest four-window multiple-instance result on 400 official test videos achieved 83.75 percent "
        "accuracy, 91.41 percent precision, 74.50 percent recall, and 93.00 percent specificity. The checkpoint "
        "remains a research artifact because accuracy, precision, and recall did not all exceed 95 percent.",
    )
    doc.add_heading("Garbage Detection", level=2)
    add_body(
        doc,
        "The combined TACO and MJU candidate achieved 67.54 percent precision, 44.49 percent recall, 49.04 percent "
        "mAP50, and 35.40 percent mAP50-95 on the 225-image untouched TACO test split. The earlier TACO-only "
        "candidate also failed, so neither expanded candidate replaced the runtime baseline.",
    )
    doc.add_heading("Fallen Person Detection", level=2)
    add_body(
        doc,
        "The expanded classifier achieved 100 percent frame-level accuracy, precision, recall, and specificity "
        "across 163 untouched frames with 45 true positives, 118 true negatives, no false positives, and no false "
        "negatives. It also achieved 100 percent on nine held-out source sequences. This candidate passed and was "
        "exported to models/fallen_person/best.pt. Testing on additional sites and camera viewpoints is still "
        "required before broad deployment because the evaluation source is one controlled dataset.",
    )
    doc.add_heading("Mobile Phone Detection", level=2)
    add_body(
        doc,
        "The expanded detector achieved 87.31 percent precision, 85.10 percent recall, 89.22 percent mAP50, and "
        "76.52 percent mAP50-95 across 793 untouched images. The result improves on the historical baseline but "
        "does not satisfy the promotion rule, so the existing runtime checkpoint remains unchanged.",
    )
    doc.add_heading("Accident Detection", level=2)
    add_body(
        doc,
        "The expanded classifier achieved 79.42 percent image-level accuracy, 79.04 percent precision, 80.09 "
        "percent recall, and 78.76 percent specificity across 452 mosaics. Aggregation across 226 held-out source "
        "videos achieved 80.09 percent accuracy, 79.31 percent precision, 81.42 percent recall, and 78.76 percent "
        "specificity. The candidate was not exported.",
    )

    doc.add_page_break()
    doc.add_heading("System Verification", level=1)
    add_table(
        doc,
        ["Verification area", "Evidence", "Result"],
        [
            [
                "Automated tests",
                "18 Pytest tests across backend, ML contracts, and frontend contract",
                "Passed",
            ],
            [
                "Python quality",
                "Ruff check across backend, machine_learning, tests, and scripts",
                "Passed",
            ],
            ["Frontend syntax", "Node syntax check for FrontEnd/script.js", "Passed"],
            [
                "Runtime models",
                "Five checkpoints loaded through the production inference adapter",
                "Passed",
            ],
            [
                "API",
                "Health, models, cameras, events, analytics, settings, frontend, infer, snapshot, stream",
                "Passed",
            ],
            [
                "Error handling",
                "Unknown model request returned 422; unsupported media returned 415",
                "Passed",
            ],
            ["Browser", "Five pages, camera Start and Stop, thresholds, empty states, console", "Passed"],
            [
                "Physical camera",
                "OpenCV camera 0, 1920 by 1080 capture, JPEG snapshot, MJPEG stream",
                "Passed with FPS limit",
            ],
        ],
        [1.5, 4.55, 0.85],
    )
    doc.add_heading("Runtime Architecture", level=2)
    for item in (
        "FastAPI serves the dashboard and versioned /api/v1 routes from one process.",
        "One OpenCV worker owns each camera source, preventing duplicate reads by browser viewers.",
        "One analysis thread per active camera sends normalized predictions to the event service.",
        "SQLite stores detections, acknowledgement state, analytics inputs, and persisted settings.",
        "The browser receives an annotated MJPEG stream and reads health, events, "
        "analytics, and configuration from the API.",
    ):
        add_bullet(doc, item)

    doc.add_heading("Camera and Performance Evidence", level=2)
    add_body(
        doc,
        "The software requests 1920 by 1080 capture and 60 FPS, with JPEG quality 92 and a separate 60 FPS stream "
        "target. The installed camera driver reports 60 FPS but the physical device sustained about 30 FPS alone "
        "and 18.5 to 20.0 FPS during simultaneous five-model inference and browser streaming. The dashboard shows "
        "the measured FPS. A true 60 FPS source requires a camera and capture interface that can deliver 60 unique "
        "1080p frames each second.",
    )
    add_body(
        doc,
        "A warmed five-model inference pass measured 110.5 ms during final API camera testing on the RTX 5050 "
        "Laptop GPU. Inference runs at a configurable 4 FPS by default so capture and display remain responsive.",
    )

    doc.add_page_break()
    doc.add_heading("Readiness and Next Engineering Work", level=1)
    add_body(
        doc,
        "The application is runnable and fully integrated. The backend, frontend, database, camera workflow, and "
        "five-model adapter passed final verification. The strict quality policy prevents failed research models "
        "from being labelled production-ready.",
    )
    add_table(
        doc,
        ["Area", "Current position", "Required before wider deployment"],
        [
            [
                "Fallen person",
                "Strict held-out gate passed",
                "Test across additional sites, viewpoints, lighting, and camera "
                "hardware",
            ],
            [
                "Fight",
                "Research result retained",
                "Add diverse temporal video data and independent cross-dataset "
                "evaluation",
            ],
            [
                "Garbage",
                "Research result retained",
                "Improve small-object labels, class coverage, and hard-negative "
                "examples",
            ],
            [
                "Mobile phone",
                "Expanded model improved but failed gate",
                "Add distant, occluded, low-light, and non-phone hand-object "
                "examples",
            ],
            [
                "Accident",
                "Temporal classifier failed gate",
                "Use longer event context, stronger motion features, and an "
                "independent video benchmark",
            ],
            [
                "Camera",
                "1080p verified; device below 60 unique FPS",
                "Use verified 1080p60 hardware or lower the displayed target to "
                "the device capability",
            ],
        ],
        [1.25, 2.05, 3.6],
    )
    doc.add_heading("Run and Review", level=2)
    add_body(
        doc,
        "Run powershell -ExecutionPolicy Bypass -File .\\scripts\\start.ps1 from the One Eye folder, then open "
        "http://127.0.0.1:8000. The dashboard exposes system health, live feed controls, detections, analytics, "
        "and confidence thresholds. Technical evidence is stored in documentation/IMPLEMENTATION_PROGRESS.md, "
        "and machine-readable final metrics and promotion decisions are stored in "
        "documentation/MODEL_EVALUATION_RESULTS.json. The latest completed pipeline stage "
        "for each model is stored in machine_learning/training_status.json.",
    )

    add_footer(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
